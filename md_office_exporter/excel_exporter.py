"""Markdown table to Excel export."""

from __future__ import annotations

from pathlib import Path
import re
from typing import Any, Mapping

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .config import load_config
from .markdown_parser import MarkdownTable, extract_markdown_tables


_INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


def _sheet_name(title: str, used_names: set[str], table_number: int) -> str:
    base = _INVALID_SHEET_CHARS.sub("_", title).strip().strip("'")
    base = base or f"Table_{table_number}"
    base = base[:31]
    candidate = base
    suffix_number = 2
    while candidate.casefold() in used_names:
        suffix = f"_{suffix_number}"
        candidate = f"{base[: 31 - len(suffix)]}{suffix}"
        suffix_number += 1
    used_names.add(candidate.casefold())
    return candidate


def _apply_sheet_style(
    worksheet: Any,
    table: MarkdownTable,
    config: Mapping[str, Any],
) -> None:
    excel_config = config["excel"]
    header_config = excel_config["header"]
    body_config = excel_config["body"]
    thin = Side(style="thin", color=excel_config["border_color"])
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    alignment = Alignment(
        horizontal=excel_config["alignment"]["horizontal"],
        vertical=excel_config["alignment"]["vertical"],
        wrap_text=excel_config["alignment"]["wrap_text"],
    )

    for row_index in range(1, table.row_count + 1):
        worksheet.row_dimensions[row_index].height = excel_config["row_height"]
        for col_index in range(1, table.column_count + 1):
            cell = worksheet.cell(row=row_index, column=col_index)
            is_header = row_index - 1 in table.header_rows
            cell.font = Font(
                name=excel_config["font"],
                size=excel_config["font_size_pt"],
                color=(
                    header_config["font_color"]
                    if is_header
                    else body_config["font_color"]
                ),
                bold=header_config["bold"] if is_header else False,
            )
            cell.alignment = alignment
            cell.border = border
            if is_header:
                cell.fill = PatternFill(
                    fill_type="solid",
                    fgColor=header_config["fill"],
                )

    for col_index in range(1, table.column_count + 1):
        max_length = 0
        for row_index in range(1, table.row_count + 1):
            value = worksheet.cell(row=row_index, column=col_index).value
            if value is not None:
                line_length = max((len(line) for line in str(value).splitlines()), default=0)
                max_length = max(max_length, line_length)
        width = min(
            excel_config["max_column_width"],
            max(excel_config["min_column_width"], max_length + 2),
        )
        worksheet.column_dimensions[get_column_letter(col_index)].width = width

    worksheet.sheet_view.showGridLines = not excel_config["hide_gridlines"]
    if excel_config["freeze_header"] and table.header_rows:
        worksheet.freeze_panes = f"A{max(table.header_rows) + 2}"
    if excel_config["auto_filter"] and table.header_rows:
        filter_header_row = max(table.header_rows) + 1
        worksheet.auto_filter.ref = (
            f"A{filter_header_row}:"
            f"{get_column_letter(table.column_count)}{table.row_count}"
        )


def _write_table(
    worksheet: Any,
    table: MarkdownTable,
    merge_mode: str,
) -> None:
    anchors = {(cell.row, cell.col): cell for cell in table.cells}
    for row_index in range(table.row_count):
        for col_index in range(table.column_count):
            if merge_mode == "expand":
                value = table.rows[row_index][col_index]
            else:
                anchor = anchors.get((row_index, col_index))
                value = anchor.value if anchor else None
            worksheet.cell(row=row_index + 1, column=col_index + 1, value=value)

    if merge_mode == "merge":
        for anchor in table.cells:
            if anchor.rowspan > 1 or anchor.colspan > 1:
                worksheet.merge_cells(
                    start_row=anchor.row + 1,
                    start_column=anchor.col + 1,
                    end_row=anchor.row + anchor.rowspan,
                    end_column=anchor.col + anchor.colspan,
                )


def export_markdown_tables_to_excel(
    markdown_path: str | Path,
    output_path: str | Path,
    merge_mode: str | None = None,
    config_path: str | Path | None = None,
    config: Mapping[str, Any] | None = None,
) -> Path:
    """Extract every Markdown/HTML table into a separate Excel worksheet."""

    source_path = Path(markdown_path)
    if not source_path.is_file():
        raise FileNotFoundError(f"Markdown 文件不存在: {source_path}")
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    resolved_config = load_config(config_path, config)
    resolved_mode = merge_mode or resolved_config["excel"]["merge_mode"]
    if resolved_mode not in {"merge", "expand"}:
        raise ValueError("merge_mode 只能是 'merge' 或 'expand'")

    tables = extract_markdown_tables(source_path.read_text(encoding="utf-8"))
    if not tables:
        raise ValueError("Markdown 文件中未找到表格")

    workbook = Workbook()
    workbook.remove(workbook.active)
    used_names: set[str] = set()
    for number, table in enumerate(tables, start=1):
        requested_name = (
            f"Sheet{number}"
            if resolved_config["excel"]["sheet_name_mode"] == "sequential"
            else table.title
        )
        worksheet = workbook.create_sheet(_sheet_name(requested_name, used_names, number))
        _write_table(worksheet, table, resolved_mode)
        _apply_sheet_style(worksheet, table, resolved_config)
        worksheet.sheet_properties.pageSetUpPr.fitToPage = True
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0

    workbook.save(output)
    return output.resolve()
