"""Markdown to DOCX export."""

from __future__ import annotations

from pathlib import Path
import re
from typing import Any, Mapping

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from .config import load_config
from .markdown_parser import MarkdownTable, parse_markdown_blocks


_INLINE_PATTERN = re.compile(
    r"(!?\[[^\]]*]\([^)]+\)|\*\*.+?\*\*|__.+?__|`[^`]+`|\*[^*]+\*|_[^_]+_)",
    flags=re.DOTALL,
)


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#").upper())


def _set_run_font(run: Any, name: str, size_pt: float | None = None) -> None:
    run.font.name = name
    if run._element.get_or_add_rPr().rFonts is None:
        run._element.get_or_add_rPr().append(OxmlElement("w:rFonts"))
    fonts = run._element.get_or_add_rPr().rFonts
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), name)
    if size_pt is not None:
        run.font.size = Pt(size_pt)


def _alignment(value: str) -> WD_ALIGN_PARAGRAPH:
    return {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }.get(value.lower(), WD_ALIGN_PARAGRAPH.LEFT)


def _set_cell_fill(cell: Any, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), color.lstrip("#").upper())


def _set_cell_margins(cell: Any, margins: Mapping[str, int]) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    name_map = {"left": "start", "right": "end", "top": "top", "bottom": "bottom"}
    for source_name, target_name in name_map.items():
        node = tc_mar.find(qn(f"w:{target_name}"))
        if node is None:
            node = OxmlElement(f"w:{target_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margins[source_name]))
        node.set(qn("w:type"), "dxa")


def _set_repeat_table_header(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    table_header = OxmlElement("w:tblHeader")
    table_header.set(qn("w:val"), "true")
    tr_pr.append(table_header)


def _set_table_borders(table: Any, color: str) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color.lstrip("#").upper())


def _set_table_fixed_layout(
    table: Any,
    widths_dxa: list[int],
    indent_dxa: int,
) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    tbl_width = tbl_pr.first_child_found_in("w:tblW")
    if tbl_width is None:
        tbl_width = OxmlElement("w:tblW")
        tbl_pr.append(tbl_width)
    tbl_width.set(qn("w:type"), "dxa")
    tbl_width.set(qn("w:w"), str(sum(widths_dxa)))

    tbl_indent = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_indent is None:
        tbl_indent = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_indent)
    tbl_indent.set(qn("w:type"), "dxa")
    tbl_indent.set(qn("w:w"), str(indent_dxa))

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_width = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            if tc_width is None:
                tc_width = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_width)
            tc_width.set(qn("w:type"), "dxa")
            tc_width.set(qn("w:w"), str(widths_dxa[min(index, len(widths_dxa) - 1)]))


def _column_widths(table_data: MarkdownTable, total_dxa: int) -> list[int]:
    scores: list[int] = []
    for col in range(table_data.column_count):
        length = max(
            (max((len(part) for part in row[col].splitlines()), default=0)
             if col < len(row) else 0)
            for row in table_data.rows
        )
        scores.append(min(40, max(8, length)))
    total_score = sum(scores) or len(scores) or 1
    widths = [max(720, round(total_dxa * score / total_score)) for score in scores]
    difference = total_dxa - sum(widths)
    if widths:
        widths[-1] += difference
    return widths


def _add_hyperlink(paragraph: Any, text: str, url: str, font_name: str) -> None:
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run_element = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    fonts = OxmlElement("w:rFonts")
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), font_name)
    run_properties.extend([fonts, color, underline])
    text_element = OxmlElement("w:t")
    text_element.text = text
    run_element.extend([run_properties, text_element])
    hyperlink.append(run_element)
    paragraph._p.append(hyperlink)


def _add_inline_text(
    paragraph: Any,
    text: str,
    font_name: str,
    base_color: str = "000000",
) -> None:
    position = 0
    for match in _INLINE_PATTERN.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position : match.start()])
            _set_run_font(run, font_name)
            run.font.color.rgb = _rgb(base_color)

        token = match.group(0)
        image_or_link = re.fullmatch(r"(!?)\[([^\]]*)]\(([^)]+)\)", token)
        if image_or_link:
            is_image, label, target = image_or_link.groups()
            if is_image:
                run = paragraph.add_run(f"[图片: {label or target}]")
                _set_run_font(run, font_name)
                run.italic = True
            else:
                _add_hyperlink(paragraph, label or target, target, font_name)
        elif token.startswith(("**", "__")):
            run = paragraph.add_run(token[2:-2])
            _set_run_font(run, font_name)
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            _set_run_font(run, "Consolas")
            run.font.color.rgb = _rgb(base_color)
        else:
            run = paragraph.add_run(token[1:-1])
            _set_run_font(run, font_name)
            run.italic = True
        position = match.end()

    if position < len(text):
        run = paragraph.add_run(text[position:])
        _set_run_font(run, font_name)
        run.font.color.rgb = _rgb(base_color)


def _configure_document(document: Document, config: Mapping[str, Any]) -> None:
    docx_config = config["docx"]
    page = docx_config["page"]
    section = document.sections[0]
    page_size = page["size"].lower()
    if page_size == "a4":
        width, height = Cm(21.0), Cm(29.7)
    elif page_size == "letter":
        width, height = Inches(8.5), Inches(11)
    else:
        raise ValueError("docx.page.size 仅支持 A4 或 Letter")
    if page["orientation"].lower() == "landscape":
        section.orientation = WD_ORIENT.LANDSCAPE
        width, height = height, width
    section.page_width, section.page_height = width, height
    margins = page["margins_cm"]
    section.top_margin = Cm(margins["top"])
    section.right_margin = Cm(margins["right"])
    section.bottom_margin = Cm(margins["bottom"])
    section.left_margin = Cm(margins["left"])

    body = docx_config["body"]
    normal = document.styles["Normal"]
    normal.font.name = body["font"]
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), body["font"])
    normal.font.size = Pt(body["font_size_pt"])
    normal.font.color.rgb = _rgb(body["color"])
    normal.paragraph_format.alignment = _alignment(body["alignment"])
    normal.paragraph_format.space_before = Pt(body["space_before_pt"])
    normal.paragraph_format.space_after = Pt(body["space_after_pt"])
    normal.paragraph_format.line_spacing = body["line_spacing"]

    for level in range(1, 7):
        heading_config = docx_config["headings"].get(
            str(level), docx_config["headings"].get("3")
        )
        style = document.styles[f"Heading {level}"]
        style.font.name = heading_config["font"]
        style._element.rPr.rFonts.set(qn("w:eastAsia"), heading_config["font"])
        style.font.size = Pt(heading_config["font_size_pt"])
        style.font.color.rgb = _rgb(heading_config["color"])
        style.font.bold = heading_config.get("bold", True)
        style.paragraph_format.space_before = Pt(heading_config["space_before_pt"])
        style.paragraph_format.space_after = Pt(heading_config["space_after_pt"])
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = document.styles[style_name]
        style.font.name = body["font"]
        style._element.rPr.rFonts.set(qn("w:eastAsia"), body["font"])
        style.font.size = Pt(body["font_size_pt"])
        style.paragraph_format.space_after = Pt(
            docx_config["lists"]["space_after_pt"]
        )

    header_text = docx_config["header"]["text"]
    if header_text:
        paragraph = section.header.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = paragraph.add_run(header_text)
        _set_run_font(run, body["font"], docx_config["header"]["font_size_pt"])
        run.font.color.rgb = _rgb(docx_config["header"]["color"])

    footer_config = docx_config["footer"]
    if footer_config["text"] or footer_config["show_page_number"]:
        paragraph = section.footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if footer_config["text"]:
            run = paragraph.add_run(footer_config["text"])
            _set_run_font(run, body["font"], footer_config["font_size_pt"])
            run.font.color.rgb = _rgb(footer_config["color"])
        if footer_config["show_page_number"]:
            if footer_config["text"]:
                paragraph.add_run("  ")
            field_begin = OxmlElement("w:fldChar")
            field_begin.set(qn("w:fldCharType"), "begin")
            instruction = OxmlElement("w:instrText")
            instruction.set(qn("xml:space"), "preserve")
            instruction.text = "PAGE"
            field_end = OxmlElement("w:fldChar")
            field_end.set(qn("w:fldCharType"), "end")
            run = paragraph.add_run()
            run._r.extend([field_begin, instruction, field_end])


def _add_table(document: Document, data: MarkdownTable, config: Mapping[str, Any]) -> None:
    if not data.row_count or not data.column_count:
        return
    table_config = config["docx"]["table"]
    table = document.add_table(rows=data.row_count, cols=data.column_count)
    table.style = table_config["style"]
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    section = document.sections[-1]
    content_width_dxa = int(
        (section.page_width - section.left_margin - section.right_margin) / 635
    )
    _set_table_fixed_layout(
        table,
        _column_widths(data, content_width_dxa),
        table_config["indent_dxa"],
    )
    _set_table_borders(table, table_config["border_color"])

    for row_index in range(data.row_count):
        if row_index in data.header_rows and table_config["repeat_header"]:
            _set_repeat_table_header(table.rows[row_index])
        for col_index in range(data.column_count):
            cell = table.cell(row_index, col_index)
            cell.text = ""
            cell.vertical_alignment = {
                "top": WD_CELL_VERTICAL_ALIGNMENT.TOP,
                "center": WD_CELL_VERTICAL_ALIGNMENT.CENTER,
                "bottom": WD_CELL_VERTICAL_ALIGNMENT.BOTTOM,
            }.get(
                table_config["vertical_alignment"].lower(),
                WD_CELL_VERTICAL_ALIGNMENT.CENTER,
            )
            _set_cell_margins(cell, table_config["cell_margin_dxa"])

    for anchor in data.cells:
        if anchor.rowspan > 1 or anchor.colspan > 1:
            start = table.cell(anchor.row, anchor.col)
            end = table.cell(
                anchor.row + anchor.rowspan - 1,
                anchor.col + anchor.colspan - 1,
            )
            start.merge(end)

    # Populate after merging. python-docx concatenates all source-cell
    # paragraphs during merge, so writing first would leave trailing blanks.
    for anchor in data.cells:
        cell = table.cell(anchor.row, anchor.col)
        cell.text = ""
        paragraph = cell.paragraphs[0]
        _add_inline_text(
            paragraph,
            anchor.value,
            table_config["font"],
            table_config["header_font_color"]
            if anchor.is_header or anchor.row in data.header_rows
            else "000000",
        )
        for run in paragraph.runs:
            _set_run_font(
                run,
                table_config["font"],
                table_config["font_size_pt"],
            )
            if anchor.is_header or anchor.row in data.header_rows:
                run.bold = True
        if anchor.is_header or anchor.row in data.header_rows:
            _set_cell_fill(cell, table_config["header_fill"])


def convert_markdown_to_docx(
    markdown_path: str | Path,
    output_path: str | Path,
    config_path: str | Path | None = None,
    config: Mapping[str, Any] | None = None,
    tables_only: bool = False,
) -> Path:
    """Convert a UTF-8 Markdown file to DOCX and return the output path.

    With ``tables_only=True`` only table blocks are exported; an error is
    raised when the Markdown contains no table at all.
    """

    source_path = Path(markdown_path)
    if not source_path.is_file():
        raise FileNotFoundError(f"Markdown 文件不存在: {source_path}")
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    resolved_config = load_config(config_path, config)
    markdown_text = source_path.read_text(encoding="utf-8")

    document = Document()
    _configure_document(document, resolved_config)
    body_font = resolved_config["docx"]["body"]["font"]

    table_count = 0
    for block in parse_markdown_blocks(markdown_text):
        block_type = block["type"]
        if tables_only and block_type != "table":
            continue
        if block_type == "heading":
            paragraph = document.add_heading(level=block["level"])
            heading_config = resolved_config["docx"]["headings"].get(
                str(block["level"]),
                resolved_config["docx"]["headings"]["3"],
            )
            _add_inline_text(
                paragraph,
                block["text"],
                heading_config["font"],
                heading_config["color"],
            )
        elif block_type == "paragraph":
            paragraph = document.add_paragraph()
            _add_inline_text(paragraph, block["text"], body_font)
        elif block_type == "list":
            style = "List Number" if block["ordered"] else "List Bullet"
            for item in block["items"]:
                paragraph = document.add_paragraph(style=style)
                indent = resolved_config["docx"]["lists"]["left_indent_cm"]
                paragraph.paragraph_format.left_indent = Cm(
                    indent + item["level"] * 0.5
                )
                paragraph.paragraph_format.first_line_indent = Cm(
                    -resolved_config["docx"]["lists"]["hanging_cm"]
                )
                _add_inline_text(paragraph, item["text"], body_font)
        elif block_type == "blockquote":
            quote_config = resolved_config["docx"]["blockquote"]
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(
                quote_config["left_indent_cm"]
            )
            _add_inline_text(
                paragraph,
                block["text"],
                body_font,
                quote_config["color"],
            )
            for run in paragraph.runs:
                run.italic = quote_config["italic"]
        elif block_type == "code":
            code_config = resolved_config["docx"]["code"]
            paragraph = document.add_paragraph()
            _set_cell_fill_like_paragraph(paragraph, code_config["fill"])
            run = paragraph.add_run(block["text"])
            _set_run_font(run, code_config["font"], code_config["font_size_pt"])
            run.font.color.rgb = _rgb(code_config["color"])
        elif block_type == "table":
            table_count += 1
            _add_table(document, block["table"], resolved_config)
        elif block_type == "horizontal_rule":
            paragraph = document.add_paragraph()
            paragraph_format = paragraph._p.get_or_add_pPr()
            borders = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "1")
            bottom.set(qn("w:color"), "B7B7B7")
            borders.append(bottom)
            paragraph_format.append(borders)

    if tables_only and table_count == 0:
        raise ValueError("Markdown 文件中未找到表格")

    document.save(output)
    return output.resolve()


def _set_cell_fill_like_paragraph(paragraph: Any, color: str) -> None:
    paragraph_properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color.lstrip("#").upper())
    paragraph_properties.append(shading)
