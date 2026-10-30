from __future__ import annotations

from pathlib import Path
import tempfile
import unittest

from docx import Document
from openpyxl import load_workbook

from md_office_exporter import (
    convert_markdown_to_docx,
    export_markdown_tables_to_excel,
    extract_markdown_tables,
)
from md_office_exporter.config import load_config


SAMPLE = """# 示例

## 普通表格

| 姓名 | 分数 |
|---|---:|
| 张三 | 95 |

## 合并表格

<table>
<thead>
<tr><th rowspan="2">部门</th><th colspan="2">季度</th></tr>
<tr><th>Q1</th><th>Q2</th></tr>
</thead>
<tbody>
<tr><td rowspan="2">研发</td><td>10</td><td>12</td></tr>
<tr><td colspan="2">合计 22</td></tr>
</tbody>
</table>
"""


class ExporterTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.markdown = self.root / "sample.md"
        self.markdown.write_text(SAMPLE, encoding="utf-8")

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_extracts_gfm_and_html_tables(self) -> None:
        tables = extract_markdown_tables(SAMPLE)
        self.assertEqual(2, len(tables))
        self.assertEqual("普通表格", tables[0].title)
        self.assertEqual("合并表格", tables[1].title)
        self.assertEqual((4, 3), (tables[1].row_count, tables[1].column_count))
        spans = {
            (cell.row, cell.col, cell.rowspan, cell.colspan)
            for cell in tables[1].cells
        }
        self.assertIn((0, 0, 2, 1), spans)
        self.assertIn((0, 1, 1, 2), spans)
        self.assertIn((2, 0, 2, 1), spans)
        self.assertIn((3, 1, 1, 2), spans)

    def test_excel_merge_mode(self) -> None:
        output = self.root / "merge.xlsx"
        export_markdown_tables_to_excel(self.markdown, output, merge_mode="merge")
        workbook = load_workbook(output)
        self.assertEqual(["普通表格", "合并表格"], workbook.sheetnames)
        sheet = workbook["合并表格"]
        merged = {str(cell_range) for cell_range in sheet.merged_cells.ranges}
        self.assertEqual({"A1:A2", "B1:C1", "A3:A4", "B4:C4"}, merged)
        self.assertEqual("季度", sheet["B1"].value)

    def test_excel_sequential_sheet_names(self) -> None:
        output = self.root / "sequential.xlsx"
        export_markdown_tables_to_excel(
            self.markdown,
            output,
            config={"excel": {"sheet_name_mode": "sequential"}},
        )
        workbook = load_workbook(output)
        self.assertEqual(["Sheet1", "Sheet2"], workbook.sheetnames)

    def test_excel_rejects_invalid_sheet_name_mode(self) -> None:
        with self.assertRaisesRegex(ValueError, "sheet_name_mode"):
            load_config(overrides={"excel": {"sheet_name_mode": "invalid"}})

    def test_excel_expand_mode_fills_covered_cells(self) -> None:
        output = self.root / "expand.xlsx"
        export_markdown_tables_to_excel(self.markdown, output, merge_mode="expand")
        workbook = load_workbook(output)
        sheet = workbook["合并表格"]
        self.assertFalse(sheet.merged_cells.ranges)
        self.assertEqual("部门", sheet["A1"].value)
        self.assertEqual("部门", sheet["A2"].value)
        self.assertEqual("季度", sheet["B1"].value)
        self.assertEqual("季度", sheet["C1"].value)
        self.assertEqual("研发", sheet["A3"].value)
        self.assertEqual("研发", sheet["A4"].value)
        self.assertEqual("合计 22", sheet["B4"].value)
        self.assertEqual("合计 22", sheet["C4"].value)

    def test_docx_export(self) -> None:
        output = self.root / "sample.docx"
        convert_markdown_to_docx(
            self.markdown,
            output,
            config={
                "docx": {
                    "body": {"font": "Arial"},
                    "headings": {"1": {"font": "Arial", "color": "C00000"}},
                }
            },
        )
        document = Document(output)
        self.assertEqual(2, len(document.tables))
        self.assertEqual("示例", document.paragraphs[0].text)
        self.assertEqual("C00000", str(document.paragraphs[0].runs[0].font.color.rgb))
        self.assertEqual("部门", document.tables[1].cell(0, 0).text)

    def test_docx_tables_only_skips_non_table_blocks(self) -> None:
        output = self.root / "tables_only.docx"
        convert_markdown_to_docx(self.markdown, output, tables_only=True)
        document = Document(output)
        self.assertEqual(2, len(document.tables))
        self.assertEqual([], [paragraph.text for paragraph in document.paragraphs])
        self.assertEqual("姓名", document.tables[0].cell(0, 0).text)
        self.assertEqual("部门", document.tables[1].cell(0, 0).text)

    def test_docx_tables_only_rejects_markdown_without_tables(self) -> None:
        plain = self.root / "plain.md"
        plain.write_text("# 只有标题\n\n一段正文，没有表格。", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "未找到表格"):
            convert_markdown_to_docx(plain, self.root / "plain.docx", tables_only=True)


if __name__ == "__main__":
    unittest.main()
